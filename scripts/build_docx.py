#!/usr/bin/env python3
"""docx 组装器：把翻译结果 spec(JSON) + 图片组装成单栏中文 docx。

用法:
  python3 build_docx.py <spec.json> <output.docx>

spec.json 结构:
{
  "meta": {
    "title_zh": "中文标题",
    "title_en": "Original English Title",
    "authors": "A. Author, B. Author",
    "venue": "IEEE Transactions on Medical Imaging (IEEE TMI)",   // 期刊/会议名
    "tier": "中科院一区 TOP / JCR Q1 (CCF-B)",                      // 等级分区
    "pub_date": "2024 年 3 月",
    "doi": "10.1109/TMI.2024.xxxxx",
    "translated_note": "由 AI 翻译归档，图表取自原文",                  // 可选
    "key_problem": "这篇论文主要解决什么问题（2-4 句，显示在信息栏下方导读区）",   // 可选但推荐
    "key_innovations": ["创新点1", "创新点2", "创新点3"]             // 可选但推荐，逐条列出
  },
  "blocks": [
    {"type": "h1", "text": "1 引言"},
    {"type": "h2", "text": "1.1 研究背景"},
    {"type": "p", "text": "正文段落，支持 **加粗** 与 *斜体*。"},
    {"type": "image", "path": "figs/fig1.png", "caption": "图1 系统总体框架", "width_cm": 15},
    {"type": "table", "caption": "表1 各方法对比结果",
     "rows": [["方法", "Dice", "HD"], ["U-Net", "0.82", "4.1"]],
     "header": true},
    {"type": "note", "text": "译注：……"}      // 小字灰色说明
  ]
}
"""
import json
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

MAX_IMG_CM = 15.0


def set_run_font(run, size=None, bold=None, italic=None, color=None,
                 east="宋体", latin="Times New Roman"):
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|\*.+?\*)")


def add_rich_text(p, text, size=11, east="宋体"):
    """解析 **加粗** 和 *斜体*。"""
    for seg in INLINE_RE.split(text):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**"):
            r = p.add_run(seg[2:-2])
            set_run_font(r, size=size, bold=True, east=east)
        elif seg.startswith("*") and seg.endswith("*") and len(seg) > 2:
            r = p.add_run(seg[1:-1])
            set_run_font(r, size=size, italic=True, east=east)
        else:
            r = p.add_run(seg)
            set_run_font(r, size=size, east=east)


def setup_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)  # A4
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(11)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    return doc


def add_meta_header(doc, meta):
    # 中文标题
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_rich_text(p, meta.get("title_zh", ""), size=16)
    p.runs and set_run_font(p.runs[0], size=16, bold=True, east="黑体")

    if meta.get("title_en"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_rich_text(p, meta["title_en"], size=12)
        if p.runs:
            set_run_font(p.runs[0], size=12, italic=True)

    if meta.get("authors"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_rich_text(p, meta["authors"], size=10.5)
        if p.runs:
            set_run_font(p.runs[0], size=10.5)

    # 出处信息表: 期刊/等级/时间/DOI
    rows = []
    if meta.get("venue"):
        rows.append(("收录期刊 / 会议", meta["venue"]))
    if meta.get("tier"):
        rows.append(("等级与分区", meta["tier"]))
    if meta.get("pub_date"):
        rows.append(("发表时间", meta["pub_date"]))
    if meta.get("doi"):
        rows.append(("DOI", meta["doi"]))
    if meta.get("translated_note"):
        rows.append(("归档说明", meta["translated_note"]))

    if rows:
        t = doc.add_table(rows=len(rows), cols=2)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, (k, v) in enumerate(rows):
            c0, c1 = t.rows[i].cells
            p0 = c0.paragraphs[0]
            add_rich_text(p0, f"**{k}**", size=10)
            if p0.runs:
                set_run_font(p0.runs[0], size=10, east="黑体")
            p1 = c1.paragraphs[0]
            add_rich_text(p1, str(v), size=10)
        t.columns[0].width = Cm(4.0)
        for row in t.rows:
            row.cells[0].width = Cm(4.0)
            row.cells[1].width = Cm(12.0)
        doc.add_paragraph()

    # 导读: 本文要解决的问题 + 核心创新点(置于信息栏与正文之间)
    if meta.get("key_problem") or meta.get("key_innovations"):
        p = doc.add_paragraph()
        add_rich_text(p, "导读", size=13)
        if p.runs:
            set_run_font(p.runs[0], size=13, bold=True, east="黑体")
        p.paragraph_format.space_after = Pt(4)
    if meta.get("key_problem"):
        p = doc.add_paragraph()
        add_rich_text(p, "**本文要解决的问题：**" + meta["key_problem"], size=10.5)
        p.paragraph_format.line_spacing = 1.25
    if meta.get("key_innovations"):
        p = doc.add_paragraph()
        add_rich_text(p, "**核心创新点：**", size=10.5)
        if p.runs:
            set_run_font(p.runs[0], size=10.5, east="黑体")
        for item in meta["key_innovations"]:
            p = doc.add_paragraph()
            add_rich_text(p, "• " + item, size=10.5)
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.line_spacing = 1.25
        doc.add_paragraph()


def add_image(doc, spec):
    path = spec["path"]
    if not os.path.exists(path):
        doc.add_paragraph(f"[缺图: {path}]")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    width = Cm(min(spec.get("width_cm", MAX_IMG_CM), MAX_IMG_CM))
    run.add_picture(path, width=width)
    if spec.get("caption"):
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_rich_text(cap, spec["caption"], size=9.5)
        if cap.runs:
            set_run_font(cap.runs[0], size=9.5, east="黑体")


def add_table(doc, spec):
    if spec.get("caption"):
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_rich_text(cap, spec["caption"], size=9.5)
        if cap.runs:
            set_run_font(cap.runs[0], size=9.5, east="黑体")
    rows = spec["rows"]
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    header = spec.get("header", True)
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            p = t.rows[i].cells[j].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_rich_text(p, str(cell), size=9.5)
            if header and i == 0 and p.runs:
                set_run_font(p.runs[0], size=9.5, bold=True, east="黑体")
    doc.add_paragraph()


def build(spec_path, out_path):
    with open(spec_path, encoding="utf-8") as f:
        spec = json.load(f)
    doc = setup_document()
    add_meta_header(doc, spec.get("meta", {}))

    for b in spec.get("blocks", []):
        t = b["type"]
        if t == "h1":
            p = doc.add_paragraph()
            add_rich_text(p, b["text"], size=14)
            if p.runs:
                set_run_font(p.runs[0], size=14, bold=True, east="黑体")
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
        elif t == "h2":
            p = doc.add_paragraph()
            add_rich_text(p, b["text"], size=12)
            if p.runs:
                set_run_font(p.runs[0], size=12, bold=True, east="黑体")
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
        elif t == "h3":
            p = doc.add_paragraph()
            add_rich_text(p, b["text"], size=11)
            if p.runs:
                set_run_font(p.runs[0], size=11, bold=True, east="黑体")
        elif t == "p":
            p = doc.add_paragraph()
            add_rich_text(p, b["text"], size=11)
            p.paragraph_format.first_line_indent = Pt(22)  # 首行缩进两字符
            p.paragraph_format.line_spacing = 1.3
        elif t == "image":
            add_image(doc, b)
        elif t == "table":
            add_table(doc, b)
        elif t == "note":
            p = doc.add_paragraph()
            add_rich_text(p, b["text"], size=9)
            if p.runs:
                set_run_font(p.runs[0], size=9, color=RGBColor(0x66, 0x66, 0x66))
        elif t == "reference":
            p = doc.add_paragraph()
            add_rich_text(p, b["text"], size=9.5)
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.75)  # 悬挂缩进
            p.paragraph_format.line_spacing = 1.2
        else:
            raise ValueError(f"未知 block 类型: {t}")

    doc.save(out_path)
    print(f"已生成: {out_path}")


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
